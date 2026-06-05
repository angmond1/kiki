# -*- coding: utf-8 -*-
"""
ki-dining 회의록 엑셀 헬퍼.

표준: `{yymmdd}_회의록.xlsx` (yymmdd = 지급신청 처리일).
같은날 처리하는 회의비 건은 모두 동일 파일에 행 추가.

9컬럼 형식:
  순번 | 사용일자 | 금액 | 장소(거래처) | 처리계정 | 내부참석자 | 외부참석자 | 외부참석자 소속 | 회의목적

회의목적 셀에는 "제목\n상세내용" 형태로 작성 (10만원 ↑ 상세 필수).

사용 예:
    from meeting_log_xlsx import open_or_create, append_row, read_log
    path = open_or_create("2026.06", "260605")
    append_row(path, {
        "date_text": "4월 30일 13:00~14:30",
        "amount": 323000,
        "place": "현화림",
        "acccd": "26E0331",
        "int_members": "이동기",
        "ext_members": "김동희, 김해진, 윤동호, 서원빈, 박정호, 윤태준, 김보민",
        "ext_org": "서울대학교",
        "title": "e-epoxidation 반응성 향상 논의",
        "content": "1. e-epoxidation 반응성 향상을 위한 촉매 개발 현황 공유\n - ...",
    })
    rows = read_log(path)   # list[dict]
"""
from __future__ import annotations
import os
from typing import Optional

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None
    # pip install openpyxl

HEADERS = ["순번", "사용일자", "금액", "장소(거래처)", "처리계정",
           "내부참석자", "외부참석자", "외부참석자 소속", "회의목적"]
WIDTHS = [6, 20, 11, 18, 10, 12, 26, 18, 55]

# 기본 루트 (사용자별로 다를 수 있음, config 로 오버라이드)
DEFAULT_ROOT = r"D:\GoogleDrive\내 드라이브\01\06.명세서"


def expected_path(year_month: str, yymmdd: str,
                  root: Optional[str] = None) -> str:
    """`{root}\{YYYY.MM}\{yymmdd}_회의록.xlsx` 경로 반환."""
    root = root or DEFAULT_ROOT
    folder = os.path.join(root, year_month)
    return os.path.join(folder, f"{yymmdd}_회의록.xlsx")


def open_or_create(year_month: str, yymmdd: str,
                   root: Optional[str] = None) -> str:
    """엑셀 파일을 열거나(있으면) 9컬럼 양식으로 생성(없으면). 경로 반환."""
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed; run: pip install openpyxl")
    fp = expected_path(year_month, yymmdd, root)
    folder = os.path.dirname(fp)
    os.makedirs(folder, exist_ok=True)
    if os.path.exists(fp):
        return fp

    wb = Workbook()
    ws = wb.active
    ws.title = "회의록"
    ws.append(HEADERS)
    for c in range(1, len(HEADERS) + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(fp)
    return fp


def _next_seq(ws) -> int:
    """현재 시트의 다음 순번."""
    seq = 0
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        v = row[0]
        if isinstance(v, (int, float)):
            seq = max(seq, int(v))
    return seq + 1


def append_row(path: str, data: dict) -> int:
    """1건 추가. data 키:
        date_text   사용일자 (예 '4월 30일 13:00~14:30')
        amount      금액 (int)
        place       장소(거래처)
        acccd       처리계정 (예 '26E0331')
        int_members 내부참석자 (성명, 콤마 구분)
        ext_members 외부참석자 (성명, 콤마 구분)
        ext_org     외부참석자 소속
        title       회의제목
        content     회의내용 (10만원↑ 상세)
    return: 부여된 순번
    """
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed")
    wb = openpyxl.load_workbook(path)
    ws = wb["회의록"] if "회의록" in wb.sheetnames else wb.active
    seq = _next_seq(ws)

    title = data.get("title", "")
    content = data.get("content", "")
    purpose_cell = (f"{title}\n{content}" if content else title)

    ws.append([
        seq,
        data.get("date_text", ""),
        int(data.get("amount", 0)),
        data.get("place", ""),
        data.get("acccd", ""),
        data.get("int_members", ""),
        data.get("ext_members", ""),
        data.get("ext_org", ""),
        purpose_cell,
    ])
    r = ws.max_row
    # 셀 포맷
    ws.cell(row=r, column=3).number_format = "#,##0"
    ws.cell(row=r, column=7).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=r, column=8).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=r, column=9).alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    return seq


def read_log(path: str) -> list[dict]:
    """엑셀 회의록 전체 행을 dict 리스트로 반환 (Claude 가 fam_0704 작성 시 조회)."""
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["회의록"] if "회의록" in wb.sheetnames else wb.active
    out = []
    for row in ws.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True):
        if not row or row[0] is None:
            continue
        purpose = row[8] or ""
        # 회의목적 셀 분리 (제목 첫 줄 / 상세 이후)
        if "\n" in (purpose or ""):
            title, content = purpose.split("\n", 1)
        else:
            title, content = purpose, ""
        out.append({
            "seq": row[0],
            "date_text": row[1],
            "amount": row[2],
            "place": row[3],
            "acccd": row[4],
            "int_members": row[5],
            "ext_members": row[6],
            "ext_org": row[7],
            "title": title,
            "content": content,
        })
    return out


if __name__ == "__main__":
    # 간단 테스트: 표본 파일 읽기
    import sys
    if len(sys.argv) > 1:
        fp = sys.argv[1]
        for r in read_log(fp):
            print(r)
