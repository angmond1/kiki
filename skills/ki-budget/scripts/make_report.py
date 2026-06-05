# -*- coding: utf-8 -*-
"""
ki-budget 예산 리포트 렌더러.
입력: fetch 수집 JSON (snapshot_date / track_categories / projects{acccd:{name,pi,role,direct{A,D},categories:{cat:{A,completed,pending,D}}}})
출력: 과제 행 x 카테고리(총액/잔액) + (한 칸 띄우고) 직접비(잔액/총액) 엑셀.
사용: python make_report.py <input.json> <output.xlsx>
credential-free. 개인 식별자/경로 하드코딩 없음(전부 인자/JSON).
"""
import json
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="999999")
LIGHT = Side(style="thin", color="CCCCCC")


def h_main(c):
    c.font = Font(bold=True, color="FFFFFF", size=11)
    c.fill = PatternFill("solid", fgColor="305496")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = Border(THIN, THIN, THIN, THIN)


def h_sub(c, color="D9E1F2"):
    c.font = Font(bold=True, size=10)
    c.fill = PatternFill("solid", fgColor=color)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = Border(THIN, THIN, THIN, THIN)


def money(c, bold=False, muted=False):
    c.number_format = '#,##0'
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.border = Border(LIGHT, LIGHT, LIGHT, LIGHT)
    c.font = Font(bold=bold, size=11, color="808080" if muted else "000000")


def txt(c, bold=False, align="left"):
    c.font = Font(bold=bold, size=11)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    c.border = Border(LIGHT, LIGHT, LIGHT, LIGHT)


def dash(ws, row, col, label="-"):
    for k in range(2):
        c = ws.cell(row=row, column=col + k, value=label)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.font = Font(color="BBBBBB")
        c.border = Border(LIGHT, LIGHT, LIGHT, LIGHT)


def main(json_path, out_path):
    snap = json.load(open(json_path, encoding="utf-8"))
    cats = snap["track_categories"]
    ds = snap["snapshot_date"]

    wb = Workbook()
    ws = wb.active
    ws.title = "예산 현황"
    ws["A1"] = f"과제별 예산 현황 — 기준일 {ds}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.row_dimensions[1].height = 24

    r1, r2 = 3, 4
    for i, name in enumerate(["과제번호", "과제명", "PI/역할"], start=1):
        c = ws.cell(row=r1, column=i, value=name)
        h_main(c)
        ws.merge_cells(start_row=r1, start_column=i, end_row=r2, end_column=i)

    col = 4
    for cat in cats:
        m = ws.cell(row=r1, column=col, value=cat)
        h_main(m)
        ws.merge_cells(start_row=r1, start_column=col, end_row=r1, end_column=col + 1)
        h_sub(ws.cell(row=r2, column=col, value="총액"))
        h_sub(ws.cell(row=r2, column=col + 1, value="잔액"), color="FCE4D6")
        col += 2

    # 학생인건비(마지막 카테고리)와 한 칸 띄우고 직접비(잔액/총액)
    blank_col = col          # 구분용 빈 열
    direct_col = col + 1     # 직접비 시작 열
    dm = ws.cell(row=r1, column=direct_col, value="직접비")
    h_main(dm)
    ws.merge_cells(start_row=r1, start_column=direct_col, end_row=r1, end_column=direct_col + 1)
    h_sub(ws.cell(row=r2, column=direct_col, value="잔액"), color="FCE4D6")
    h_sub(ws.cell(row=r2, column=direct_col + 1, value="총액"))

    ws.row_dimensions[r1].height = 22
    ws.row_dimensions[r2].height = 20

    row = r2 + 1
    for code, p in snap["projects"].items():
        txt(ws.cell(row=row, column=1, value=code), bold=True, align="center")
        txt(ws.cell(row=row, column=2, value=p.get("name", "")), bold=True)
        txt(ws.cell(row=row, column=3, value=f'{p.get("pi","")}/{p.get("role","")}'), align="center")
        col = 4
        for cat in cats:
            data = p.get("categories", {}).get(cat)
            if not data:
                dash(ws, row, col)
            else:
                money(ws.cell(row=row, column=col, value=data.get("A", 0)), muted=True)
                money(ws.cell(row=row, column=col + 1, value=data.get("D", 0)), bold=True)
            col += 2
        # 직접비 (빈 열 건너뛰고 잔액/총액)
        direct = p.get("direct")
        if direct:
            money(ws.cell(row=row, column=direct_col, value=direct.get("D", 0)), bold=True)
            money(ws.cell(row=row, column=direct_col + 1, value=direct.get("A", 0)), muted=True)
        ws.row_dimensions[row].height = 22
        row += 1

    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 12
    col = 4
    for _ in cats:
        ws.column_dimensions[get_column_letter(col)].width = 14
        ws.column_dimensions[get_column_letter(col + 1)].width = 14
        col += 2
    ws.column_dimensions[get_column_letter(blank_col)].width = 3      # 구분 빈 열(좁게)
    ws.column_dimensions[get_column_letter(direct_col)].width = 15
    ws.column_dimensions[get_column_letter(direct_col + 1)].width = 15
    ws.freeze_panes = ws.cell(row=r2 + 1, column=4)

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"[OK] 저장: {out}")
    return out


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
