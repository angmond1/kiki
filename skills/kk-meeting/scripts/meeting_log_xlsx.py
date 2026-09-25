# -*- coding: utf-8 -*-
"""
kk-meeting 회의록 엑셀 헬퍼.

표준(2026-09-24~): `{root}\{yymm}_회의록.xlsx` — 한 폴더, 월별 1파일(yymm = 지급신청 처리 연월). 그 달 처리 건은 모두 같은 파일에 행 추가.
주된 목적 = 이전 회의 주제·내용과의 중복 방지 기록(지급신청에 첨부하지 않음). 새 회의록 전 all_titles() 로 전부 스캔.

9컬럼 형식:
  순번 | 사용일자 | 금액 | 장소(거래처) | 처리계정 | 내부참석자 | 외부참석자 | 외부참석자 소속 | 회의목적

회의목적 셀에는 "제목\n상세내용" 형태로 작성 (10만원 ↑ 상세 필수).

사용 예:
    from meeting_log_xlsx import open_or_create, append_row, read_log
    path = open_or_create("2606")            # {root}\2606_회의록.xlsx
    append_row(path, {
        "date_text": "4월 30일 13:00~14:30",
        "amount": 323000,
        "place": "○○식당",
        "acccd": "26E0001",
        "int_members": "홍길동",
        "ext_members": "김철수, 이영희, 박민수, 정지원, 최유리, 강현우, 윤서연",
        "ext_org": "○○대학교",
        "title": "연구 진행상황 논의",
        "content": "1. 연구 진행상황 및 향후 계획 공유\n - ...",
    })
    rows = read_log(path)   # list[dict]
    past = all_titles()      # 중복 방지: 폴더 내 모든 회의록의 제목·내용
"""
from __future__ import annotations
import os
import re
from typing import Optional

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None
    # pip install openpyxl

HEADERS = ["순번", "사용일자", "금액", "장소(거래처)", "처리계정",
           "내부참석자", "외부참석자", "외부참석자 소속", "회의목적"]
WIDTHS = [6, 20, 11, 18, 10, 12, 26, 18, 55]

# 기본 루트 (사용자별로 다를 수 있음, config 로 오버라이드)
DEFAULT_ROOT = r"C:\kiki\meeting\meeting_log"


def expected_path(yymm: str, root: Optional[str] = None) -> str:
    """`{root}\{yymm}_회의록.xlsx` 경로 반환 (한 폴더, 월별 1파일)."""
    root = root or DEFAULT_ROOT
    return os.path.join(root, f"{yymm}_회의록.xlsx")


def open_or_create(yymm: str, root: Optional[str] = None) -> str:
    """엑셀 파일을 열거나(있으면) 9컬럼 양식으로 생성(없으면). 경로 반환."""
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed; run: pip install openpyxl")
    fp = expected_path(yymm, root)
    folder = os.path.dirname(fp)
    os.makedirs(folder, exist_ok=True)
    if os.path.exists(fp):
        return fp

    wb = Workbook()
    ws = wb.active
    ws.title = "회의록"
    ws.append(HEADERS)
    for c in range(1, len(HEADERS) + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(fp)
    return fp


def _to_amount(v) -> int:
    """323000 / '323,000' / '45,000
11,000'(같은 날 식당+카페 두 금액) → 정수 합계. 숫자 없으면 0."""
    if isinstance(v, (int, float)):
        return int(v)
    nums = re.findall(r"\d[\d,]*", str(v or ""))
    return sum(int(n.replace(",", "")) for n in nums) if nums else 0


def _next_seq(ws) -> int:
    """현재 시트의 다음 순번."""
    seq = 0
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        v = row[0]
        if isinstance(v, (int, float)):
            seq = max(seq, int(v))
    return seq + 1


def append_row(path: str, data: dict) -> int:
    """1건 추가. data 키:
        date_text   사용일자 (예 '4월 30일 13:00~14:30')
        amount      금액 (int)
        place       장소(거래처)
        acccd       처리계정 (예 '26E0001')
        int_members 내부참석자 (성명, 콤마 구분)
        ext_members 외부참석자 (성명, 콤마 구분)
        ext_org     외부참석자 소속
        title       회의제목
        content     회의내용 (10만원↑ 상세)
    return: 부여된 순번
    """
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed")
    wb = openpyxl.load_workbook(path)
    ws = wb["회의록"] if "회의록" in wb.sheetnames else wb.active
    seq = _next_seq(ws)

    title = data.get("title", "")
    content = data.get("content", "")
    purpose_cell = (f"{title}\n{content}" if content else title)

    ws.append([
        seq,
        data.get("date_text", ""),
        _to_amount(data.get("amount", 0)),
        data.get("place", ""),
        data.get("acccd", ""),
        data.get("int_members", ""),
        data.get("ext_members", ""),
        data.get("ext_org", ""),
        purpose_cell,
    ])
    r = ws.max_row
    # 셀 포맷
    ws.cell(row=r, column=3).number_format = "#,##0"
    ws.cell(row=r, column=7).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=r, column=8).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=r, column=9).alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    return seq


def read_log(path: str) -> list[dict]:
    """엑셀 회의록 전체 행을 dict 리스트로 반환 (Claude 가 fam_0704 작성 시 조회)."""
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["회의록"] if "회의록" in wb.sheetnames else wb.active
    out = []
    for row in ws.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True):
        if not row or row[0] is None:
            continue
        purpose = row[8] or ""
        # 회의목적 셀 분리 (제목 첫 줄 / 상세 이후)
        if "\n" in (purpose or ""):
            title, content = purpose.split("\n", 1)
        else:
            title, content = purpose, ""
        out.append({
            "seq": row[0],
            "date_text": row[1],
            "amount": row[2],
            "place": row[3],
            "acccd": row[4],
            "int_members": row[5],
            "ext_members": row[6],
            "ext_org": row[7],
            "title": title,
            "content": content,
        })
    return out


def all_titles(root: Optional[str] = None) -> list[dict]:
    """중복 방지용 — 폴더의 모든 `*_회의록.xlsx` 를 읽어
    [{file, date_text, amount, place, acccd, title, content}] 를 반환한다.
    새 회의록을 쓰기 전 사용자가 준 주제를 title 들과 비교(같거나 유사하면 조정 제안)."""
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed")
    root = root or DEFAULT_ROOT
    out: list[dict] = []
    if not os.path.isdir(root):
        return out
    for fn in sorted(os.listdir(root)):
        if not fn.endswith("_회의록.xlsx") or fn.startswith("~$"):
            continue
        try:
            ws = openpyxl.load_workbook(os.path.join(root, fn), data_only=True).active
        except Exception:
            continue
        for r in ws.iter_rows(min_row=2, values_only=True):
            r = list(r) + [None] * 9
            if not any(r[:9]):
                continue
            purpose = str(r[8] or "")
            title, _, content = purpose.partition("\n")
            out.append({"file": fn, "date_text": str(r[1] or ""), "amount": r[2], "place": str(r[3] or ""),
                        "acccd": str(r[4] or ""), "title": title.strip(), "content": content.strip()})
    return out


if __name__ == "__main__":
    # 간단 테스트: 표본 파일 읽기
    import sys
    if len(sys.argv) > 1:
        fp = sys.argv[1]
        for r in read_log(fp):
            print(r)
